import unittest
import pandas as pd
from customer_profitability.synthetic import build_synthetic_dataset
from customer_profitability.analysis import analyze_profitability
from customer_profitability.sensitivity import run_sensitivity
from customer_profitability.validation import DataValidationError
from customer_profitability.reporting import REPORT_SHEETS, generate_demo_artifacts
from openpyxl import load_workbook
import tempfile

class ProfitabilityModelTest(unittest.TestCase):
    def setUp(self): self.dataset = build_synthetic_dataset()
    def test_dataset_is_complete_and_repeatable(self):
        other = build_synthetic_dataset(); self.assertEqual(len(self.dataset.transactions), 240); pd.testing.assert_frame_equal(self.dataset.transactions, other.transactions)
    def test_profit_waterfall_reconciles(self):
        result = analyze_profitability(self.dataset.transactions); expected = result.detail["net_revenue"] - result.detail["product_cost"] - result.detail["fulfillment_cost"] - result.detail["service_cost"]; self.assertLess((expected - result.detail["contribution_profit"]).abs().max(), 1e-8)
    def test_customer_and_product_summaries_reconcile(self):
        result = analyze_profitability(self.dataset.transactions); total = result.metrics["contribution_profit"]; self.assertAlmostEqual(result.customer_summary["contribution_profit"].sum(), total); self.assertAlmostEqual(result.product_summary["contribution_profit"].sum(), total)
    def test_high_revenue_low_profit_customer_is_identified(self):
        flagged = analyze_profitability(self.dataset.transactions).customer_summary.query("profitability_flag == '高收入低利润'"); self.assertGreaterEqual(len(flagged), 1)
    def test_profitability_threshold_is_configurable(self):
        strict = analyze_profitability(self.dataset.transactions, low_margin_threshold=0.50)
        default = analyze_profitability(self.dataset.transactions)
        self.assertGreaterEqual((strict.customer_summary["profitability_flag"] == "高收入低利润").sum(), (default.customer_summary["profitability_flag"] == "高收入低利润").sum())
    def test_invalid_profitability_threshold_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "阈值"):
            analyze_profitability(self.dataset.transactions, low_margin_threshold=1.1)
    def test_hhi_is_bounded(self):
        hhi = analyze_profitability(self.dataset.transactions).metrics["revenue_hhi"]; self.assertGreater(hhi, 0); self.assertLessEqual(hhi, 1)
    def test_sensitivity_direction(self):
        result = analyze_profitability(self.dataset.transactions); scenarios = run_sensitivity(result, self.dataset.scenarios).set_index("scenario"); self.assertGreater(scenarios.loc["提价优化", "contribution_profit"], scenarios.loc["基准", "contribution_profit"]); self.assertLess(scenarios.loc["降价承压", "contribution_profit"], scenarios.loc["基准", "contribution_profit"])
    def test_duplicate_transaction_is_rejected(self):
        duplicate = pd.concat([self.dataset.transactions, self.dataset.transactions.iloc[[0]]], ignore_index=True)
        with self.assertRaisesRegex(DataValidationError, "不能重复"): analyze_profitability(duplicate)
    def test_discount_and_returns_cannot_exceed_revenue(self):
        invalid = self.dataset.transactions.copy(); invalid.loc[0, "discount"] = invalid.loc[0, "list_revenue"]; invalid.loc[0, "returns"] = 1
        with self.assertRaisesRegex(DataValidationError, "不能超过"): analyze_profitability(invalid)
    def test_demo_workbooks_reopen_and_contain_reports(self):
        with tempfile.TemporaryDirectory() as temporary:
            artifacts=generate_demo_artifacts(temporary); workbook=load_workbook(artifacts.report_path)
            self.assertEqual(workbook.sheetnames,REPORT_SHEETS); self.assertEqual(len(workbook["客户盈利"]._charts),1); self.assertEqual(len(workbook["敏感性"]._charts),1); workbook.close()
    def test_empty_transactions_are_rejected(self):
        with self.assertRaisesRegex(DataValidationError,"不能为空"): analyze_profitability(self.dataset.transactions.iloc[0:0])
    def test_excel_numeric_text_and_month_are_normalized(self):
        frame=self.dataset.transactions.copy(); frame["quantity"]=frame["quantity"].astype(str); frame["month"]=frame["month"].dt.strftime("%Y-%m-%d"); result=analyze_profitability(frame); self.assertEqual(len(result.detail),240)
    def test_invalid_scenario_factor_is_rejected(self):
        result=analyze_profitability(self.dataset.transactions); invalid=self.dataset.scenarios.copy(); invalid.loc[0,"price_change"]=-1
        with self.assertRaisesRegex(ValueError,"大于 -100%"): run_sensitivity(result,invalid)
    def test_input_workbook_is_labeled_synthetic(self):
        with tempfile.TemporaryDirectory() as temporary:
            artifacts=generate_demo_artifacts(temporary); workbook=load_workbook(artifacts.input_path,read_only=True); messages=[r[0] for r in workbook["演示说明"].iter_rows(min_row=2,values_only=True)]; self.assertTrue(any("合成数据" in str(x) for x in messages)); workbook.close()

if __name__ == "__main__": unittest.main()
